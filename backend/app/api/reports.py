
from datetime import date
from decimal import Decimal
from enum import Enum
from io import BytesIO
import re

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import inspect as sa_inspect, select

import app.models as models
from app.core.dependencies import DbSession, HRUser


router = APIRouter(
    prefix="/reports",
    tags=["Reports & Exports"],
)

MAX_REPORT_ROWS = 5000

REPORT_MODELS = {
    "employees": ["Employee"],
    "attendance": ["Attendance"],
    "payroll": ["Payroll"],
}


def mapped_model(name: str):
    """Return an existing mapped model, if it exists."""
    candidate = getattr(models, name, None)

    if candidate is None:
        return None

    mapper = sa_inspect(candidate, raiseerr=False)

    if mapper is None or not hasattr(mapper, "column_attrs"):
        return None

    return candidate


def model_columns(model):
    mapper = sa_inspect(model)
    return [column.key for column in mapper.column_attrs]


def discover_attrition_model():
    """
    Find an existing prediction/history model.
    Do not create or modify any database tables.
    """
    candidates = []

    for candidate in vars(models).values():
        if not isinstance(candidate, type):
            continue

        mapper = sa_inspect(candidate, raiseerr=False)

        if mapper is None or not hasattr(mapper, "column_attrs"):
            continue

        name = (
            candidate.__name__ + " " + mapper.local_table.name
        ).lower()

        columns = {key.lower() for key in model_columns(candidate)}

        if not any(
            word in name
            for word in ("attrition", "prediction", "risk")
        ):
            continue

        prediction_fields = {
            "risk_score",
            "risk_percentage",
            "risk_level",
            "attrition_probability",
            "prediction",
            "predicted_attrition",
            "confidence_score",
            "probability",
        }

        matching_fields = columns.intersection(prediction_fields)

        if not matching_fields:
            continue

        score = len(matching_fields)

        if "prediction" in name:
            score += 3

        if "history" in name:
            score += 3

        if "attrition" in name:
            score += 2

        if "alert" in name:
            score -= 2

        candidates.append((score, candidate))

    if not candidates:
        return None

    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0][1]


def json_value(value):
    """Convert database values into export-friendly values."""
    if isinstance(value, Enum):
        return value.value

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, (date,)):
        return value.isoformat()

    if value is None or isinstance(value, (str, int, float, bool)):
        return value

    if hasattr(value, "isoformat"):
        return value.isoformat()

    return str(value)


def find_date_column(model):
    columns = model_columns(model)

    preferred = (
        "attendance_date",
        "payroll_date",
        "pay_period",
        "prediction_date",
        "created_at",
        "joining_date",
        "date",
    )

    for candidate in preferred:
        if candidate in columns:
            return getattr(model, candidate)

    for name in columns:
        if name.endswith("_date") or name.endswith("_at"):
            return getattr(model, name)

    return None


def fetch_report_rows(
    db,
    report_type: str,
    date_from: date | None,
    date_to: date | None,
):
    if report_type == "attrition":
        model = discover_attrition_model()

        if model is None:
            raise HTTPException(
                status_code=501,
                detail=(
                    "No suitable attrition prediction/history model "
                    "was found in app.models. Connect this report to "
                    "your existing attrition history model."
                ),
            )
    else:
        model = None

        for name in REPORT_MODELS[report_type]:
            model = mapped_model(name)
            if model is not None:
                break

        if model is None:
            raise HTTPException(
                status_code=500,
                detail=(
                    f"The existing SQLAlchemy model for "
                    f"{report_type} could not be found."
                ),
            )

    statement = select(model)

    if date_from or date_to:
        date_column = find_date_column(model)

        if date_column is None:
            raise HTTPException(
                status_code=422,
                detail=(
                    f"No date column was found for {report_type}. "
                    "Check the model's date field before filtering."
                ),
            )

        if date_from:
            statement = statement.where(date_column >= date_from)

        if date_to:
            statement = statement.where(date_column <= date_to)

    columns = model_columns(model)

    if "id" in columns:
        statement = statement.order_by(getattr(model, "id").desc())

    statement = statement.limit(MAX_REPORT_ROWS)

    records = db.scalars(statement).all()

    # Employee reports also include the related User information
    # without exposing password hashes or other authentication fields.
    rows = []

    for record in records:
        row = {
            column: json_value(getattr(record, column))
            for column in columns
        }

        if report_type == "employees":
            user = getattr(record, "user", None)

            if user is not None:
                for field in (
                    "email",
                    "first_name",
                    "last_name",
                    "is_active",
                ):
                    if hasattr(user, field):
                        row[f"user_{field}"] = json_value(
                            getattr(user, field)
                        )

            department = getattr(record, "department", None)

            if department is not None and hasattr(department, "name"):
                row["department_name"] = department.name

        rows.append(row)

    headers = list(dict.fromkeys(
        key
        for row in rows
        for key in row
    ))

    if not headers:
        headers = columns

    if report_type == "employees" and not rows:
        headers = columns + [
            "user_email",
            "user_first_name",
            "user_last_name",
            "user_is_active",
            "department_name",
        ]

    return headers, rows


def create_excel(title, headers, rows):
    output = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"

    sheet.append([title])
    sheet.append([])

    sheet.append(headers)

    for cell in sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1D4ED8",
        )

    for row in rows:
        sheet.append([
            row.get(header) for header in headers
        ])

    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = (
        f"A3:{sheet.cell(3, len(headers)).column_letter}"
        f"{max(sheet.max_row, 3)}"
    )

    for column_cells in sheet.columns:
        letter = column_cells[0].column_letter
        max_length = max(
            (len(str(cell.value or "")) for cell in column_cells),
            default=10,
        )
        sheet.column_dimensions[letter].width = min(
            max(max_length + 2, 12),
            35,
        )

    workbook.save(output)
    return output.getvalue()


def create_pdf(title, headers, rows):
    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24,
    )

    styles = getSampleStyleSheet()
    content = [
        Paragraph(title, styles["Title"]),
        Spacer(1, 12),
    ]

    if not headers:
        headers = ["Message"]

    table_data = [
        [Paragraph(str(header), styles["BodyText"]) for header in headers]
    ]

    for row in rows:
        table_data.append([
            Paragraph(
                str(row.get(header, "") if row.get(header) is not None else "")
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;"),
                styles["BodyText"],
            )
            for header in headers
        ])

    if len(table_data) == 1:
        table_data.append(["No records found"] + [""] * (len(headers) - 1))

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))

    content.append(table)
    document.build(content)

    return output.getvalue()


@router.get("/{report_type}.{file_format}")
def export_report(
    report_type: str,
    file_format: str,
    db: DbSession,
    current_user: HRUser,
    date_from: date | None = Query(default=None),
    date_to: date | None = Query(default=None),
):
    """
    Export employee, attrition, payroll or attendance reports.
    Supported formats: pdf and xlsx.
    """

    allowed_reports = {
        "employees",
        "attrition",
        "payroll",
        "attendance",
    }

    if report_type not in allowed_reports:
        raise HTTPException(
            status_code=404,
            detail="Unknown report type",
        )

    if file_format not in {"pdf", "xlsx"}:
        raise HTTPException(
            status_code=422,
            detail="Supported formats are pdf and xlsx",
        )

    if date_from and date_to and date_from > date_to:
        raise HTTPException(
            status_code=422,
            detail="date_from cannot be after date_to",
        )

    headers, rows = fetch_report_rows(
        db,
        report_type,
        date_from,
        date_to,
    )

    title = f"WorkForce AI Pro - {report_type.title()} Report"

    if file_format == "xlsx":
        content = create_excel(title, headers, rows)
        media_type = (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    else:
        content = create_pdf(title, headers, rows)
        media_type = "application/pdf"

    filename = f"{report_type}_report.{file_format}"

    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Report-Row-Count": str(len(rows)),
            "X-Report-Truncated": str(
                len(rows) >= MAX_REPORT_ROWS
            ).lower(),
        },
    )
